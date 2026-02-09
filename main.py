from fastapi import FastAPI, Query, UploadFile, File, HTTPException, Request, Depends
from fastapi.staticfiles import StaticFiles
from fastapi.responses import StreamingResponse, FileResponse, JSONResponse, RedirectResponse
from sqlalchemy import create_engine, text
from urllib.parse import quote_plus
import os
from typing import Optional
from pydantic import BaseModel
from datetime import datetime
import uuid
from io import BytesIO
import pandas as pd
from openpyxl.utils import get_column_letter
from pathlib import Path
import hashlib
import secrets
import re

from starlette.middleware.sessions import SessionMiddleware

from etl_engine import process_excel_bytes

app = FastAPI()

# --- SESSION CONFIG ---
SECRET_KEY = os.getenv("SECRET_KEY", "change-this-secret")  # <-- เปลี่ยนค่าให้ยาวและเดายาก
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, same_site="lax")

# --- AUTH CONFIG ---
def _hash_password(raw_password: str) -> str:
    salt = "dashboard-salt"  # เปลี่ยนค่าให้ยาวและเดายาก
    return hashlib.sha256(f"{salt}:{raw_password}".encode("utf-8")).hexdigest()

USERS = {
    "admin": {"password_hash": _hash_password("admin123"), "role": "Admin"},
    "employee": {"password_hash": _hash_password("employee123"), "role": "Employee"},
}

def _verify_password(raw_password: str, password_hash: str) -> bool:
    return secrets.compare_digest(_hash_password(raw_password), password_hash)

def _normalize_username(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    text_value = str(value).strip().lower()
    return text_value or None

def get_current_user(request: Request):
    user = request.session.get("user")
    if not user:
        raise HTTPException(status_code=401, detail="unauthorized")
    return user

def get_current_user_optional(request: Request):
    return request.session.get("user")

def require_admin(user=Depends(get_current_user)):
    if user.get("role") != "Admin":
        raise HTTPException(status_code=403, detail="forbidden")
    return user

# --- CONFIG ---
db_password = quote_plus("teezaza123") # <--- 🔑 แก้รหัสผ่านของคุณ
default_db_url = f'postgresql://postgres:{db_password}@localhost:5432/safety_db'
DB_CONNECTION_STR = os.getenv("DATABASE_URL", default_db_url)
engine = create_engine(DB_CONNECTION_STR)

def init_employee_table():
    create_sql = """
        CREATE TABLE IF NOT EXISTS employees (
            id SERIAL PRIMARY KEY,
            username VARCHAR(120),
            password_hash VARCHAR(255),
            area_code VARCHAR(50),
            first_name VARCHAR(120),
            last_name VARCHAR(120),
            team VARCHAR(120),
            territory VARCHAR(120),
            nickname VARCHAR(120),
            email VARCHAR(255)
        )
    """
    with engine.connect() as conn:
        conn.execute(text(create_sql))
        conn.execute(text("ALTER TABLE employees ADD COLUMN IF NOT EXISTS username VARCHAR(120)"))
        conn.execute(text("ALTER TABLE employees ADD COLUMN IF NOT EXISTS password_hash VARCHAR(255)"))
        conn.execute(text("CREATE UNIQUE INDEX IF NOT EXISTS employees_username_uq ON employees (username)"))
        conn.commit()

init_employee_table()

def init_customer_table():
    create_sql = """
        CREATE TABLE IF NOT EXISTS customers (
            id SERIAL PRIMARY KEY,
            customer_code VARCHAR(120),
            customer_name VARCHAR(200),
            province VARCHAR(120),
            region VARCHAR(120)
        )
    """
    with engine.connect() as conn:
        conn.execute(text(create_sql))
        conn.commit()

init_customer_table()

def init_sales_transactions_table():
    create_sql = """
        CREATE TABLE IF NOT EXISTS sales_transactions (
            document_date DATE,
            invoice_no VARCHAR(120),
            customer_code VARCHAR(120),
            customer_name VARCHAR(200),
            province VARCHAR(120),
            sales_rep_code VARCHAR(120),
            sales_rep_name VARCHAR(200),
            sales_team VARCHAR(120),
            product_code VARCHAR(120),
            product_group VARCHAR(120),
            product_name VARCHAR(200),
            quantity NUMERIC,
            unit_of_measure VARCHAR(50),
            unit_price NUMERIC,
            discount_percent NUMERIC,
            bill_discount_percent NUMERIC,
            unit_price_non_vat NUMERIC,
            total_amount_non_vat NUMERIC,
            batch_id VARCHAR(64)
        )
    """
    with engine.connect() as conn:
        conn.execute(text(create_sql))
        conn.execute(text("ALTER TABLE sales_transactions ADD COLUMN IF NOT EXISTS batch_id VARCHAR(64)"))
        conn.commit()

init_sales_transactions_table()

def init_update_history_table():
    create_sql = """
        CREATE TABLE IF NOT EXISTS update_history (
            id SERIAL PRIMARY KEY,
            batch_id VARCHAR(64) UNIQUE,
            source VARCHAR(50),
            filename VARCHAR(255),
            rows_count INTEGER,
            uploaded_by VARCHAR(120),
            created_at TIMESTAMP DEFAULT NOW()
        )
    """
    with engine.connect() as conn:
        conn.execute(text(create_sql))
        conn.commit()

init_update_history_table()

def init_user_profile_table():
    create_sql = """
        CREATE TABLE IF NOT EXISTS user_profiles (
            username VARCHAR(120) PRIMARY KEY,
            full_name VARCHAR(200),
            nickname VARCHAR(120),
            email VARCHAR(255),
            phone VARCHAR(50),
            team VARCHAR(120),
            territory VARCHAR(120),
            position VARCHAR(120)
        )
    """
    with engine.connect() as conn:
        conn.execute(text(create_sql))
        conn.commit()

init_user_profile_table()

PROVINCES_BY_REGION = {
    "ภาคเหนือ": [
        "เชียงใหม่", "เชียงราย", "ลำพูน", "ลำปาง", "แพร่", "น่าน", "พะเยา", "แม่ฮ่องสอน",
        "ตาก", "สุโขทัย", "พิษณุโลก", "พิจิตร", "เพชรบูรณ์", "อุตรดิตถ์", "กำแพงเพชร",
        "นครสวรรค์", "อุทัยธานี"
    ],
    "ภาคตะวันออกเฉียงเหนือ": [
        "นครราชสีมา", "บุรีรัมย์", "สุรินทร์", "ศรีสะเกษ", "อุบลราชธานี", "ยโสธร", "ชัยภูมิ",
        "อำนาจเจริญ", "บึงกาฬ", "หนองบัวลำภู", "ขอนแก่น", "อุดรธานี", "เลย", "หนองคาย",
        "มหาสารคาม", "ร้อยเอ็ด", "กาฬสินธุ์", "สกลนคร", "นครพนม", "มุกดาหาร"
    ],
    "ภาคกลาง": [
        "กรุงเทพมหานคร", "นนทบุรี", "ปทุมธานี", "พระนครศรีอยุธยา", "อ่างทอง", "ลพบุรี",
        "สิงห์บุรี", "ชัยนาท", "สระบุรี", "นครปฐม", "สมุทรสาคร", "สมุทรสงคราม",
        "สมุทรปราการ", "สุพรรณบุรี"
    ],
    "ภาคตะวันออก": [
        "ชลบุรี", "ระยอง", "จันทบุรี", "ตราด", "ฉะเชิงเทรา", "ปราจีนบุรี", "สระแก้ว", "นครนายก"
    ],
    "ภาคตะวันตก": [
        "กาญจนบุรี", "ราชบุรี", "เพชรบุรี", "ประจวบคีรีขันธ์"
    ],
    "ภาคใต้": [
        "ชุมพร", "ระนอง", "สุราษฎร์ธานี", "พังงา", "ภูเก็ต", "กระบี่", "นครศรีธรรมราช",
        "ตรัง", "พัทลุง", "สตูล", "สงขลา", "ปัตตานี", "ยะลา", "นราธิวาส"
    ]
}
REGIONS = list(PROVINCES_BY_REGION.keys())
PROVINCES = [p for region in PROVINCES_BY_REGION.values() for p in region]

# ตั้งค่าเป้าหมายยอดขายรายปี (แก้ไขตามต้องการ)
YEARLY_SALES_TARGETS = {
    2024: 0,
    2025: 0,
    2026: 0
}

TEMPLATE_COLUMNS = [
    'วันที่/เดือน/ปี เอกสาร', 'เลขที่บิล', 'รหัสลูกค้า', 'ชื่อลูกค้า', 'จังหวัด',
    'รหัสพนักงานขาย', 'ชื่อพนักงาน', 'ทีม', 'รหัสสินค้า', 'กลุ่มสินค้า', 'รายละเอียด',
    'จำนวน', 'หน่วยนับ', '@', '%ส่วนลด', '%ลดท้ายบิล', 'หน่วยละ NON VAT', 'รวมเงิน NON VAT'
]

CUSTOMER_SUMMARY_TEMPLATE_COLUMNS = [
    'วันที่เอกสาร', 'เดือน', 'ปี', 'เลขที่บิล', 'รหัสลูกค้า/ชื่อลูกค้า', 'รหัสลูกค้า',
    'ชื่อลูกค้า', 'จังหวัด', 'รหัสพนักงานขาย', 'ชือพนักงาน', 'ทีม', 'รหัสสินค้า',
    'กลุ่มสินค้า', 'รายละเอียด', 'แถม 24', 'จำนวน', 'หน่วยนับ', '@', '% ส่วนลด',
    '% ลดท้ายบิล', 'หน่วยละ NON VAT', 'รวมเงิน NON VAT'
]

class TransactionIn(BaseModel):
    document_date: str
    invoice_no: Optional[str] = None
    customer_code: Optional[str] = None
    customer_name: Optional[str] = None
    province: Optional[str] = None
    sales_rep_code: Optional[str] = None
    sales_rep_name: Optional[str] = None
    sales_team: Optional[str] = None
    product_code: Optional[str] = None
    product_group: Optional[str] = None
    product_name: Optional[str] = None
    quantity: Optional[float] = 0
    unit_of_measure: Optional[str] = None
    unit_price: Optional[float] = 0
    discount_percent: Optional[float] = 0
    bill_discount_percent: Optional[float] = 0
    unit_price_non_vat: Optional[float] = 0
    total_amount_non_vat: Optional[float] = 0

class EmployeeIn(BaseModel):
    username: Optional[str] = None
    password: Optional[str] = None
    area_code: Optional[str] = None
    first_name: Optional[str] = None
    last_name: Optional[str] = None
    team: Optional[str] = None
    territory: Optional[str] = None
    nickname: Optional[str] = None
    email: Optional[str] = None

class CustomerIn(BaseModel):
    customer_code: Optional[str] = None
    customer_name: Optional[str] = None
    province: Optional[str] = None
    region: Optional[str] = None

class ProfileIn(BaseModel):
    full_name: Optional[str] = None
    nickname: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    team: Optional[str] = None
    territory: Optional[str] = None
    position: Optional[str] = None

# ฟังก์ชันช่วยสร้าง WHERE Clause ตาม Filter ที่เลือก
def _to_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None

def build_filter(year, month, team, rep, region, province):
    conditions = []
    params = {}

    year_int = _to_int(year)
    month_int = _to_int(month) if month and month != 'All' else None
    
    # Filter บังคับ: ปี (ถ้าไม่ส่งมาจะใช้ปีปัจจุบันในการเปรียบเทียบไม่ได้ แต่ในที่นี้เราจะ filter ตอน query)
    if year_int is not None:
        conditions.append("EXTRACT(YEAR FROM document_date) = :year")
        params['year'] = year_int
    
    # Filter ทางเลือก
    if month_int is not None:
        conditions.append("EXTRACT(MONTH FROM document_date) = :month")
        params['month'] = month_int
    if team and team != 'All':
        conditions.append("sales_team = :team")
        params['team'] = team
    if rep and rep != 'All':
        conditions.append("sales_rep_name = :rep")
        params['rep'] = rep
    if province and province != 'All':
        conditions.append("province = :province")
        params['province'] = province
    if region and region != 'All':
        region_provinces = PROVINCES_BY_REGION.get(region, [])
        if region_provinces:
            placeholders = []
            for idx, name in enumerate(region_provinces):
                key = f"region_prov_{idx}"
                placeholders.append(f":{key}")
                params[key] = name
            conditions.append(f"province IN ({', '.join(placeholders)})")
        else:
            conditions.append("1 = 0")
        
    where_clause = " AND ".join(conditions)
    if where_clause:
        where_clause = "WHERE " + where_clause
    else:
        where_clause = "" # ระวัง: ถ้าไม่มี filter เลยจะดึงหมด
        
    return where_clause, params

def get_region_for_province(province_name: Optional[str]) -> Optional[str]:
    if not province_name:
        return None
    for region, provinces in PROVINCES_BY_REGION.items():
        if province_name in provinces:
            return region
    return None

def normalize_customer_code(value: Optional[object]) -> Optional[str]:
    if value is None:
        return None
    text_value = str(value).strip()
    if re.fullmatch(r"\d+\.0", text_value):
        return text_value[:-2]
    return text_value

# 1. API สำหรับตัวเลือกใน Dropdown (Filters)
@app.get("/api/options")
def get_options(user=Depends(get_current_user)):
    with engine.connect() as conn:
        years = conn.execute(text("SELECT DISTINCT EXTRACT(YEAR FROM document_date) FROM sales_transactions ORDER BY 1 DESC")).fetchall()
        teams = conn.execute(text("SELECT DISTINCT sales_team FROM sales_transactions WHERE sales_team IS NOT NULL ORDER BY 1")).fetchall()
        reps = conn.execute(text("SELECT DISTINCT sales_rep_name FROM sales_transactions WHERE sales_rep_name IS NOT NULL ORDER BY 1")).fetchall()
        
        return {
            "years": [int(row[0]) for row in years if row[0] is not None],
            "teams": [row[0] for row in teams],
            "reps": [row[0] for row in reps],
            "regions": REGIONS,
            "provinces_by_region": PROVINCES_BY_REGION,
            "provinces": PROVINCES
        }

@app.get("/api/customer_options")
def get_customer_options(user=Depends(get_current_user)):
    sql = """
        SELECT DISTINCT customer_code, customer_name
        FROM sales_transactions
        WHERE customer_code IS NOT NULL OR customer_name IS NOT NULL
        ORDER BY customer_name NULLS LAST, customer_code NULLS LAST
    """
    with engine.connect() as conn:
        rows = conn.execute(text(sql)).fetchall()

    items = []
    for row in rows:
        code = normalize_customer_code(row[0])
        name = row[1]
        if code and name:
            label = f"{name} ({code})"
            value = code
        elif code:
            label = code
            value = code
        else:
            label = name
            value = name
        items.append({"label": label, "value": value, "code": code, "name": name})

    return {"items": items}

@app.get("/api/home_summary")
def get_home_summary(user=Depends(get_current_user)):
    with engine.connect() as conn:
        sales_count = conn.execute(text("SELECT COUNT(*) FROM sales_transactions")).scalar() or 0
        customer_count = conn.execute(text("SELECT COUNT(*) FROM customers")).scalar() or 0
        employee_count = conn.execute(text("SELECT COUNT(*) FROM employees")).scalar() or 0
        latest_update = conn.execute(text("SELECT MAX(created_at) FROM update_history")).scalar()
        sparkline_rows = conn.execute(text("""
            SELECT
                document_date::date AS doc_date,
                SUM(total_amount_non_vat) AS total_amount,
                COUNT(DISTINCT COALESCE(customer_code, customer_name)) AS customer_count,
                COUNT(DISTINCT sales_rep_name) AS rep_count
            FROM sales_transactions
            WHERE document_date IS NOT NULL
            GROUP BY document_date::date
            ORDER BY document_date::date DESC
            LIMIT 12
        """)).fetchall()

    sparkline_rows = list(reversed(sparkline_rows))
    sparkline_sales = [float(row[1] or 0) for row in sparkline_rows]
    sparkline_customers = [int(row[2] or 0) for row in sparkline_rows]
    sparkline_employees = [int(row[3] or 0) for row in sparkline_rows]

    return {
        "sales_count": int(sales_count),
        "customer_count": int(customer_count),
        "employee_count": int(employee_count),
        "latest_update": latest_update.isoformat() if latest_update else None,
        "sparkline_sales": sparkline_sales,
        "sparkline_customers": sparkline_customers,
        "sparkline_employees": sparkline_employees
    }

@app.get("/api/home_feed")
def get_home_feed(user=Depends(get_current_user)):
    sql = """
        SELECT document_date, customer_name, customer_code, sales_team, total_amount_non_vat
        FROM sales_transactions
        ORDER BY document_date DESC NULLS LAST, invoice_no DESC NULLS LAST
        LIMIT 10
    """
    with engine.connect() as conn:
        rows = conn.execute(text(sql)).fetchall()

    items = []
    for row in rows:
        customer_name = row[1] or row[2] or "(ไม่ระบุลูกค้า)"
        items.append({
            "document_date": row[0].isoformat() if row[0] else None,
            "customer_name": customer_name,
            "sales_team": row[3] or "-",
            "total_amount": float(row[4] or 0)
        })

    return {"items": items}

# 1.0 API ดาวน์โหลดไฟล์ตัวอย่าง
@app.get("/api/template")
def download_template(user=Depends(require_admin)):
    df = pd.DataFrame(columns=TEMPLATE_COLUMNS)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="template")
        ws = writer.book["template"]
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22

        widths = {
            "A": 18, "B": 14, "C": 14, "D": 20, "E": 14, "F": 16, "G": 18,
            "H": 10, "I": 14, "J": 16, "K": 24, "L": 10, "M": 12, "N": 12,
            "O": 10, "P": 12, "Q": 16, "R": 18
        }
        for col in range(1, len(TEMPLATE_COLUMNS) + 1):
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].width = widths.get(col_letter, 14)
    output.seek(0)

    headers = {
        "Content-Disposition": "attachment; filename=sales_template.xlsx"
    }
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

@app.get("/api/customer_summary_template")
def download_customer_summary_template(user=Depends(require_admin)):
    df = pd.DataFrame(columns=CUSTOMER_SUMMARY_TEMPLATE_COLUMNS)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="customer_summary")
        ws = writer.book["customer_summary"]
        ws.freeze_panes = "A2"
    output.seek(0)

    headers = {
        "Content-Disposition": "attachment; filename=customer_summary_template.xlsx"
    }
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

# 1.1 API อัปโหลดไฟล์ Excel
@app.post("/api/upload_excel")
async def upload_excel(file: UploadFile = File(...), user=Depends(require_admin)):
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="รองรับเฉพาะไฟล์ Excel (.xlsx, .xls)")

    content = await file.read()
    batch_id = uuid.uuid4().hex
    result = process_excel_bytes(content, batch_id=batch_id)
    if not result.get("success"):
        raise HTTPException(status_code=400, detail=result.get("error", "นำเข้าไฟล์ไม่สำเร็จ"))

    with engine.connect() as conn:
        conn.execute(
            text("""
                INSERT INTO update_history (batch_id, source, filename, rows_count, uploaded_by)
                VALUES (:batch_id, :source, :filename, :rows_count, :uploaded_by)
            """),
            {
                "batch_id": batch_id,
                "source": "excel",
                "filename": file.filename,
                "rows_count": int(result.get("rows", 0)),
                "uploaded_by": user.get("username")
            }
        )
        conn.commit()

    return {"success": True, "rows": result.get("rows", 0), "batch_id": batch_id}

# 1.2 API เพิ่มรายการขายแบบกรอกฟอร์ม
@app.post("/api/add_transaction")
def add_transaction(payload: TransactionIn, user=Depends(require_admin)):
    try:
        doc_date = datetime.strptime(payload.document_date, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="รูปแบบวันที่ต้องเป็น YYYY-MM-DD")

    batch_id = f"manual-{uuid.uuid4().hex}"
    insert_sql = text("""
        INSERT INTO sales_transactions (
            document_date, invoice_no, customer_code, customer_name, province,
            sales_rep_code, sales_rep_name, sales_team,
            product_code, product_group, product_name,
            quantity, unit_of_measure, unit_price, discount_percent, bill_discount_percent,
            unit_price_non_vat, total_amount_non_vat, batch_id
        ) VALUES (
            :document_date, :invoice_no, :customer_code, :customer_name, :province,
            :sales_rep_code, :sales_rep_name, :sales_team,
            :product_code, :product_group, :product_name,
            :quantity, :unit_of_measure, :unit_price, :discount_percent, :bill_discount_percent,
            :unit_price_non_vat, :total_amount_non_vat, :batch_id
        )
    """)

    params = {
        "document_date": doc_date,
        "invoice_no": payload.invoice_no,
        "customer_code": payload.customer_code,
        "customer_name": payload.customer_name,
        "province": payload.province,
        "sales_rep_code": payload.sales_rep_code,
        "sales_rep_name": payload.sales_rep_name,
        "sales_team": payload.sales_team,
        "product_code": payload.product_code,
        "product_group": payload.product_group,
        "product_name": payload.product_name,
        "quantity": payload.quantity or 0,
        "unit_of_measure": payload.unit_of_measure,
        "unit_price": payload.unit_price or 0,
        "discount_percent": payload.discount_percent or 0,
        "bill_discount_percent": payload.bill_discount_percent or 0,
        "unit_price_non_vat": payload.unit_price_non_vat or 0,
        "total_amount_non_vat": payload.total_amount_non_vat or 0,
        "batch_id": batch_id
    }

    with engine.connect() as conn:
        conn.execute(insert_sql, params)
        conn.execute(
            text("""
                INSERT INTO update_history (batch_id, source, filename, rows_count, uploaded_by)
                VALUES (:batch_id, :source, :filename, :rows_count, :uploaded_by)
            """),
            {
                "batch_id": batch_id,
                "source": "manual",
                "filename": None,
                "rows_count": 1,
                "uploaded_by": user.get("username")
            }
        )
        conn.commit()

    return {"success": True, "batch_id": batch_id}

@app.get("/api/update_history")
def get_update_history(user=Depends(require_admin)):
    sql = """
        SELECT batch_id, source, filename, rows_count, uploaded_by, created_at
        FROM update_history
        ORDER BY created_at DESC
    """
    with engine.connect() as conn:
        rows = conn.execute(text(sql)).fetchall()

    return {
        "items": [
            {
                "batch_id": row[0],
                "source": row[1],
                "filename": row[2],
                "rows_count": int(row[3] or 0),
                "uploaded_by": row[4],
                "created_at": row[5].isoformat() if row[5] else None
            }
            for row in rows
        ]
    }

@app.delete("/api/update_history/{batch_id}")
def delete_update_history(batch_id: str, user=Depends(require_admin)):
    with engine.connect() as conn:
        history_row = conn.execute(
            text("SELECT 1 FROM update_history WHERE batch_id = :batch_id"),
            {"batch_id": batch_id}
        ).fetchone()
        if not history_row:
            raise HTTPException(status_code=404, detail="ไม่พบประวัติการอัปเดต")

        deleted_rows = conn.execute(
            text("DELETE FROM sales_transactions WHERE batch_id = :batch_id"),
            {"batch_id": batch_id}
        ).rowcount
        conn.execute(
            text("DELETE FROM update_history WHERE batch_id = :batch_id"),
            {"batch_id": batch_id}
        )
        conn.commit()

    return {"success": True, "deleted_rows": int(deleted_rows or 0)}

# 2. API สำหรับ KPI Cards (ยอดขาย & ยอด Shop)
@app.get("/api/kpi")
def get_kpi(
    year: int,
    month: Optional[str] = 'All',
    team: Optional[str] = 'All',
    rep: Optional[str] = 'All',
    region: Optional[str] = 'All',
    province: Optional[str] = 'All',
    user=Depends(get_current_user)
):
    where, params = build_filter(year, month, team, rep, region, province)
    
    # ถ้าเลือกเดือน -> ต้องหายอดสะสม (YTD) ถึงเดือนนั้น
    # ถ้าไม่เลือกเดือน -> ยอดสะสมคือทั้งปี
    
    sql = f"""
        SELECT 
            -- 1. ยอดขายรวม (ตาม Filter)
            COALESCE(SUM(total_amount_non_vat), 0) as sales_selected,
            -- 2. จำนวน Shop (ตาม Filter)
            COUNT(DISTINCT customer_code) as shop_selected
        FROM sales_transactions
        {where}
    """
    
    # หา YTD (สะสมตั้งแต่ต้นปี) - กรณีมีการเลือกเดือน
    ytd_where = ""
    ytd_params = params.copy()
    if month and month != 'All':
        # เปลี่ยนเงื่อนไขเดือน เป็น <= เดือนที่เลือก
        base_conditions = [c for c in where.replace("WHERE ", "").split(" AND ") if "EXTRACT(MONTH" not in c]
        base_conditions.append("EXTRACT(MONTH FROM document_date) <= :month")
        ytd_where = "WHERE " + " AND ".join(base_conditions)
    else:
        ytd_where = where # ถ้าไม่เลือกเดือน YTD ก็คือค่าเดียวกับ Selected
        
    sql_ytd = f"""
        SELECT 
            COALESCE(SUM(total_amount_non_vat), 0) as sales_ytd,
            COUNT(DISTINCT customer_code) as shop_ytd
        FROM sales_transactions
        {ytd_where}
    """

    with engine.connect() as conn:
        curr = conn.execute(text(sql), params).fetchone()
        ytd = conn.execute(text(sql_ytd), ytd_params).fetchone()
        
        return {
            "sales_period": float(curr[0]),
            "shop_period": int(curr[1]),
            "sales_accum": float(ytd[0]),
            "shop_accum": int(ytd[1]),
            "sales_target_year": float(YEARLY_SALES_TARGETS.get(year, 0))
        }

# 3. API กราฟเปรียบเทียบปี (Year vs Year)
@app.get("/api/compare_year")
def get_compare_year(
    year: int,
    team: Optional[str] = 'All',
    rep: Optional[str] = 'All',
    region: Optional[str] = 'All',
    province: Optional[str] = 'All',
    user=Depends(get_current_user)
):
    # สร้าง Filter แบบไม่เอา "ปี" และ "เดือน" (เพราะเราจะดึง 2 ปีมาเทียบกันรายเดือน)
    conditions = []
    params = {'y1': year, 'y2': year - 1}
    
    if team and team != 'All':
        conditions.append("sales_team = :team")
        params['team'] = team
    if rep and rep != 'All':
        conditions.append("sales_rep_name = :rep")
        params['rep'] = rep
    if province and province != 'All':
        conditions.append("province = :province")
        params['province'] = province
    if region and region != 'All':
        region_provinces = PROVINCES_BY_REGION.get(region, [])
        if region_provinces:
            placeholders = []
            for idx, name in enumerate(region_provinces):
                key = f"region_prov_{idx}"
                placeholders.append(f":{key}")
                params[key] = name
            conditions.append(f"province IN ({', '.join(placeholders)})")
        else:
            conditions.append("1 = 0")
        
    base_where = " AND ".join(conditions)
    if base_where: base_where = "AND " + base_where

    sql = f"""
        SELECT 
            EXTRACT(MONTH FROM document_date) as m,
            SUM(CASE WHEN EXTRACT(YEAR FROM document_date) = :y1 THEN total_amount_non_vat ELSE 0 END) as sales_current,
            SUM(CASE WHEN EXTRACT(YEAR FROM document_date) = :y2 THEN total_amount_non_vat ELSE 0 END) as sales_prev
        FROM sales_transactions
        WHERE EXTRACT(YEAR FROM document_date) IN (:y1, :y2)
        {base_where}
        GROUP BY m
        ORDER BY m
    """
    
    with engine.connect() as conn:
        result = conn.execute(text(sql), params).fetchall()
        
        # จัด Data ให้ครบ 12 เดือน (กันเหนียวเผื่อเดือนไหนไม่มีขาย)
        months = list(range(1, 13))
        data_map = {int(row[0]): (float(row[1]), float(row[2])) for row in result}
        
        return {
            "labels": ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."],
            "current_year": [data_map.get(m, (0,0))[0] for m in months],
            "prev_year": [data_map.get(m, (0,0))[1] for m in months]
        }

# 4. API Top 10 Ranking
@app.get("/api/ranking")
def get_ranking(
    year: int,
    month: Optional[str] = 'All',
    team: Optional[str] = 'All',
    rep: Optional[str] = 'All',
    region: Optional[str] = 'All',
    province: Optional[str] = 'All',
    user=Depends(get_current_user)
):
    where, params = build_filter(year, month, team, rep, region, province)
    
    with engine.connect() as conn:
        # Top 10 Products
        prod_sql = f"""
            SELECT product_name, SUM(total_amount_non_vat) as total
            FROM sales_transactions {where}
            GROUP BY product_name ORDER BY total DESC LIMIT 10
        """
        top_products = conn.execute(text(prod_sql), params).fetchall()
        
        # Top 10 Customers
        cust_sql = f"""
            SELECT customer_name, SUM(total_amount_non_vat) as total
            FROM sales_transactions {where}
            GROUP BY customer_name ORDER BY total DESC LIMIT 10
        """
        top_customers = conn.execute(text(cust_sql), params).fetchall()
        
        return {
            "products": [{"label": row[0], "value": float(row[1])} for row in top_products],
            "customers": [{"label": row[0], "value": float(row[1])} for row in top_customers]
        }

# 5. API Pie: Sales by Province
@app.get("/api/sales_by_province")
def get_sales_by_province(
    year: int,
    month: Optional[str] = 'All',
    team: Optional[str] = 'All',
    rep: Optional[str] = 'All',
    region: Optional[str] = 'All',
    province: Optional[str] = 'All',
    user=Depends(get_current_user)
):
    where, params = build_filter(year, month, team, rep, region, province)

    sql = f"""
        SELECT province, COALESCE(SUM(total_amount_non_vat), 0) as total
        FROM sales_transactions
        {where}
        GROUP BY province
        ORDER BY total DESC
    """

    with engine.connect() as conn:
        rows = conn.execute(text(sql), params).fetchall()

    items = []
    for row in rows:
        label = row[0] or "(ไม่ระบุจังหวัด)"
        items.append({"label": label, "value": float(row[1])})

    # จำกัดจำนวนชิ้นเพื่อความอ่านง่าย และรวมที่เหลือเป็น "อื่นๆ"
    max_items = 10
    if len(items) > max_items:
        head = items[:max_items]
        others_total = sum(i["value"] for i in items[max_items:])
        head.append({"label": "อื่นๆ", "value": float(others_total)})
        items = head

    return {"items": items}

# 6. API Pie YTD: Sales by Province (สะสมตั้งแต่ต้นปีถึงเดือนที่เลือก)
@app.get("/api/sales_by_province_ytd")
def get_sales_by_province_ytd(
    year: int,
    month: Optional[str] = 'All',
    team: Optional[str] = 'All',
    rep: Optional[str] = 'All',
    region: Optional[str] = 'All',
    province: Optional[str] = 'All',
    user=Depends(get_current_user)
):
    where, params = build_filter(year, month, team, rep, region, province)

    ytd_where = where
    if month and month != 'All':
        base_conditions = [c for c in where.replace("WHERE ", "").split(" AND ") if "EXTRACT(MONTH" not in c]
        base_conditions.append("EXTRACT(MONTH FROM document_date) <= :month")
        ytd_where = "WHERE " + " AND ".join(base_conditions)

    sql = f"""
        SELECT province, COALESCE(SUM(total_amount_non_vat), 0) as total
        FROM sales_transactions
        {ytd_where}
        GROUP BY province
        ORDER BY total DESC
    """

    with engine.connect() as conn:
        rows = conn.execute(text(sql), params).fetchall()

    items = []
    for row in rows:
        label = row[0] or "(ไม่ระบุจังหวัด)"
        items.append({"label": label, "value": float(row[1])})

    max_items = 10
    if len(items) > max_items:
        head = items[:max_items]
        others_total = sum(i["value"] for i in items[max_items:])
        head.append({"label": "อื่นๆ", "value": float(others_total)})
        items = head

    return {"items": items}

# 6.2 Customer purchase summary (by product & month)
@app.get("/api/customer_purchase_summary")
def get_customer_purchase_summary(
    year: int,
    customer: str = Query(..., min_length=1),
    user=Depends(get_current_user)
):
    sql = """
        SELECT
            COALESCE(product_code, '(ไม่ระบุรหัสสินค้า)') as product_code,
            COALESCE(product_name, '(ไม่ระบุชื่อสินค้า)') as product_name,
            COALESCE(unit_price, 0) as unit_price,
            EXTRACT(MONTH FROM document_date) as month,
            COALESCE(SUM(quantity), 0) as qty
        FROM sales_transactions
        WHERE EXTRACT(YEAR FROM document_date) = :year
          AND (customer_code = :customer OR customer_name = :customer)
        GROUP BY product_code, product_name, unit_price, month
        ORDER BY product_name ASC
    """

    params = {"year": year, "customer": customer}
    with engine.connect() as conn:
        rows = conn.execute(text(sql), params).fetchall()

    items_map = {}
    for row in rows:
        key = (row[0], row[1], float(row[2] or 0))
        month = int(row[3]) if row[3] is not None else None
        qty = float(row[4] or 0)
        if key not in items_map:
            items_map[key] = {
                "product_code": row[0],
                "product_name": row[1],
                "unit_price": float(row[2] or 0),
                "months": {str(m): 0 for m in range(1, 13)},
                "total": 0
            }
        if month:
            items_map[key]["months"][str(month)] += qty
            items_map[key]["total"] += qty

    return {"items": list(items_map.values())}

# 7. Employees (Admin only)
@app.get("/api/employees")
def list_employees(user=Depends(require_admin)):
    sql = """
        SELECT id, username, area_code, first_name, last_name, team, territory, nickname, email
        FROM employees
        ORDER BY id DESC
    """
    with engine.connect() as conn:
        rows = conn.execute(text(sql)).fetchall()

    return {
        "items": [
            {
                "id": row[0],
                "username": row[1],
                "area_code": row[2],
                "first_name": row[3],
                "last_name": row[4],
                "team": row[5],
                "territory": row[6],
                "nickname": row[7],
                "email": row[8]
            }
            for row in rows
        ]
    }

@app.post("/api/employees")
def add_employee(payload: EmployeeIn, user=Depends(require_admin)):
    normalized_username = _normalize_username(payload.username)
    password_hash = _hash_password(payload.password) if payload.password else None
    sql = text("""
        INSERT INTO employees (username, password_hash, area_code, first_name, last_name, team, territory, nickname, email)
        VALUES (:username, :password_hash, :area_code, :first_name, :last_name, :team, :territory, :nickname, :email)
        RETURNING id
    """)
    params = payload.dict()
    params["username"] = normalized_username
    params["password_hash"] = password_hash
    with engine.connect() as conn:
        row = conn.execute(sql, params).fetchone()
        conn.commit()
    return {"success": True, "id": row[0] if row else None}

@app.put("/api/employees/{employee_id}")
def update_employee(employee_id: int, payload: EmployeeIn, user=Depends(require_admin)):
    params = payload.dict()
    params["id"] = employee_id
    params["username"] = _normalize_username(params.get("username"))

    set_clauses = [
        "username = :username",
        "area_code = :area_code",
        "first_name = :first_name",
        "last_name = :last_name",
        "team = :team",
        "territory = :territory",
        "nickname = :nickname",
        "email = :email"
    ]

    if params.get("password"):
        params["password_hash"] = _hash_password(params["password"])
        set_clauses.append("password_hash = :password_hash")

    sql = text(f"""
        UPDATE employees
        SET {', '.join(set_clauses)}
        WHERE id = :id
    """)
    with engine.connect() as conn:
        result = conn.execute(sql, params)
        conn.commit()

    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="ไม่พบพนักงาน")

    return {"success": True}

@app.delete("/api/employees/{employee_id}")
def delete_employee(employee_id: int, user=Depends(require_admin)):
    sql = text("DELETE FROM employees WHERE id = :id")
    with engine.connect() as conn:
        result = conn.execute(sql, {"id": employee_id})
        conn.commit()

    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="ไม่พบพนักงาน")

    return {"success": True}

@app.post("/api/employees/upload")
async def upload_employees_excel(file: UploadFile = File(...), user=Depends(require_admin)):
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="รองรับเฉพาะไฟล์ Excel (.xlsx, .xls)")

    content = await file.read()
    df = pd.read_excel(BytesIO(content))
    df.columns = df.columns.str.strip().str.replace(r"\s+", " ", regex=True)

    column_mapping = {
        "Username": "username",
        "ชื่อผู้ใช้": "username",
        "Password": "password",
        "รหัสผ่าน": "password",
        "รหัสเขต": "area_code",
        "ผู้แทน": "first_name",
        "ชื่อ": "first_name",
        "นามสกุล": "last_name",
        "ทีม": "team",
        "เขต": "territory",
        "ชื่อเล่น": "nickname",
        "Mail": "email",
        "Email": "email",
        "E-mail": "email",
    }

    df = df.rename(columns=column_mapping)
    allowed_cols = ["username", "password", "area_code", "first_name", "last_name", "team", "territory", "nickname", "email"]
    df = df[[c for c in allowed_cols if c in df.columns]]
    df = df.dropna(how="all")

    if "username" in df.columns:
        df["username"] = df["username"].apply(_normalize_username)

    if "password" in df.columns:
        df["password_hash"] = df["password"].apply(lambda v: _hash_password(v) if pd.notna(v) and str(v).strip() else None)
        df = df.drop(columns=["password"])

    if df.empty:
        raise HTTPException(status_code=400, detail="ไม่พบข้อมูลที่นำเข้าได้")

    with engine.connect() as conn:
        df.to_sql("employees", engine, index=False, if_exists="append")

    return {"success": True, "rows": int(len(df))}

@app.get("/api/employees/template")
def download_employee_template(user=Depends(require_admin)):
    columns = ["Username", "Password", "รหัสเขต", "ผู้แทน", "นามสกุล", "ทีม", "เขต", "ชื่อเล่น", "Mail"]
    df = pd.DataFrame(columns=columns)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="employees")
        ws = writer.book["employees"]
        ws.freeze_panes = "A2"
    output.seek(0)

    headers = {"Content-Disposition": "attachment; filename=employees_template.xlsx"}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

# 8. Customers (Admin only)
@app.get("/api/customers")
def list_customers(
    region: Optional[str] = None,
    province: Optional[str] = None,
    search: Optional[str] = None,
    user=Depends(require_admin)
):
    conditions = []
    params = {}

    provinces_filter = []
    if province:
        provinces_filter = [p.strip() for p in province.split(",") if p.strip()]

    if provinces_filter:
        placeholders = []
        for idx, name in enumerate(provinces_filter):
            key = f"prov_{idx}"
            placeholders.append(f":{key}")
            params[key] = name
        conditions.append(f"province IN ({', '.join(placeholders)})")
    elif region:
        regions_filter = [r.strip() for r in region.split(",") if r.strip()]
        region_provinces = []
        for r in regions_filter:
            region_provinces.extend(PROVINCES_BY_REGION.get(r, []))
        if region_provinces:
            placeholders = []
            for idx, name in enumerate(region_provinces):
                key = f"regprov_{idx}"
                placeholders.append(f":{key}")
                params[key] = name
            conditions.append(f"province IN ({', '.join(placeholders)})")

    if search:
        search_value = f"%{search.strip()}%"
        conditions.append("(customer_code ILIKE :search OR customer_name ILIKE :search)")
        params["search"] = search_value

    where_clause = ""
    if conditions:
        where_clause = "WHERE " + " AND ".join(conditions)

    sql = f"""
        SELECT id, customer_code, customer_name, province, region
        FROM customers
        {where_clause}
        ORDER BY id DESC
    """
    with engine.connect() as conn:
        rows = conn.execute(text(sql), params).fetchall()

    return {
        "items": [
            {
                "id": row[0],
                "customer_code": normalize_customer_code(row[1]),
                "customer_name": row[2],
                "province": row[3],
                "region": row[4]
            }
            for row in rows
        ]
    }

@app.post("/api/customers")
def add_customer(payload: CustomerIn, user=Depends(require_admin)):
    region = payload.region or get_region_for_province(payload.province)
    sql = text("""
        INSERT INTO customers (customer_code, customer_name, province, region)
        VALUES (:customer_code, :customer_name, :province, :region)
        RETURNING id
    """)
    params = payload.dict()
    params["customer_code"] = normalize_customer_code(params.get("customer_code"))
    params["region"] = region
    with engine.connect() as conn:
        row = conn.execute(sql, params).fetchone()
        conn.commit()
    return {"success": True, "id": row[0] if row else None}

@app.put("/api/customers/{customer_id}")
def update_customer(customer_id: int, payload: CustomerIn, user=Depends(require_admin)):
    region = payload.region or get_region_for_province(payload.province)
    sql = text("""
        UPDATE customers
        SET customer_code = :customer_code,
            customer_name = :customer_name,
            province = :province,
            region = :region
        WHERE id = :id
    """)
    params = payload.dict()
    params["customer_code"] = normalize_customer_code(params.get("customer_code"))
    params["region"] = region
    params["id"] = customer_id
    with engine.connect() as conn:
        result = conn.execute(sql, params)
        conn.commit()
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="ไม่พบลูกค้า")
    return {"success": True}

@app.delete("/api/customers/{customer_id}")
def delete_customer(customer_id: int, user=Depends(require_admin)):
    sql = text("DELETE FROM customers WHERE id = :id")
    with engine.connect() as conn:
        result = conn.execute(sql, {"id": customer_id})
        conn.commit()
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="ไม่พบลูกค้า")
    return {"success": True}

@app.post("/api/customers/upload")
async def upload_customers_excel(file: UploadFile = File(...), user=Depends(require_admin)):
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="รองรับเฉพาะไฟล์ Excel (.xlsx, .xls)")

    content = await file.read()
    df = pd.read_excel(BytesIO(content))
    df.columns = df.columns.str.strip().str.replace(r"\s+", " ", regex=True)

    column_mapping = {
        "รหัสลูกค้า": "customer_code",
        "รหัสลูกค้า.1": "customer_code",
        "Customer Code": "customer_code",
        "ชื่อลูกค้า": "customer_name",
        "Customer Name": "customer_name",
        "จังหวัด": "province",
        "Province": "province"
    }

    df = df.rename(columns=column_mapping)
    allowed_cols = ["customer_code", "customer_name", "province"]
    df = df[[c for c in allowed_cols if c in df.columns]]
    df = df.dropna(how="all")

    if df.empty:
        raise HTTPException(status_code=400, detail="ไม่พบข้อมูลที่นำเข้าได้")

    if "customer_code" in df.columns:
        df["customer_code"] = df["customer_code"].apply(normalize_customer_code)
    df["region"] = df.get("province").apply(get_region_for_province)

    with engine.connect() as conn:
        df.to_sql("customers", engine, index=False, if_exists="append")

    return {"success": True, "rows": int(len(df))}

@app.get("/api/customers/template")
def download_customers_template(user=Depends(require_admin)):
    columns = ["รหัสลูกค้า", "ชื่อลูกค้า", "จังหวัด"]
    df = pd.DataFrame(columns=columns)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="customers")
        ws = writer.book["customers"]
        ws.freeze_panes = "A2"
    output.seek(0)

    headers = {"Content-Disposition": "attachment; filename=customers_template.xlsx"}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

# 8. Profile
@app.get("/api/profile")
def get_profile(user=Depends(get_current_user)):
    sql = """
        SELECT full_name, nickname, email, phone, team, territory, position
        FROM user_profiles
        WHERE username = :username
    """
    with engine.connect() as conn:
        row = conn.execute(text(sql), {"username": user.get("username")}).fetchone()

    if not row:
        return {
            "full_name": None,
            "nickname": None,
            "email": None,
            "phone": None,
            "team": None,
            "territory": None,
            "position": None
        }

    return {
        "full_name": row[0],
        "nickname": row[1],
        "email": row[2],
        "phone": row[3],
        "team": row[4],
        "territory": row[5],
        "position": row[6]
    }

@app.put("/api/profile")
def update_profile(payload: ProfileIn, user=Depends(get_current_user)):
    sql = text("""
        INSERT INTO user_profiles (username, full_name, nickname, email, phone, team, territory, position)
        VALUES (:username, :full_name, :nickname, :email, :phone, :team, :territory, :position)
        ON CONFLICT (username)
        DO UPDATE SET
            full_name = EXCLUDED.full_name,
            nickname = EXCLUDED.nickname,
            email = EXCLUDED.email,
            phone = EXCLUDED.phone,
            team = EXCLUDED.team,
            territory = EXCLUDED.territory,
            position = EXCLUDED.position
    """)
    params = payload.dict()
    params["username"] = user.get("username")

    with engine.connect() as conn:
        conn.execute(sql, params)
        conn.commit()

    return {"success": True}

# 8.1 Employee Profile (read-only for employees)
@app.get("/api/employee_profile")
def get_employee_profile(user=Depends(get_current_user)):
    if user.get("role") != "Employee":
        raise HTTPException(status_code=403, detail="forbidden")

    sql = """
        SELECT username, first_name, last_name, nickname, email, team, territory
        FROM employees
        WHERE username = :username
    """
    with engine.connect() as conn:
        row = conn.execute(text(sql), {"username": user.get("username")}).fetchone()

    if not row:
        return {
            "username": user.get("username"),
            "full_name": None,
            "nickname": None,
            "email": None,
            "team": None,
            "territory": None
        }

    full_name = " ".join([p for p in [row[1], row[2]] if p]) or None
    return {
        "username": row[0],
        "full_name": full_name,
        "nickname": row[3],
        "email": row[4],
        "team": row[5],
        "territory": row[6]
    }

@app.get("/api/team_members")
def get_team_members(user=Depends(get_current_user)):
    if user.get("role") != "Employee":
        raise HTTPException(status_code=403, detail="forbidden")

    team_sql = "SELECT team FROM employees WHERE username = :username"
    with engine.connect() as conn:
        team_row = conn.execute(text(team_sql), {"username": user.get("username")}).fetchone()
        team_name = team_row[0] if team_row else None

        if not team_name:
            return {"team": None, "items": []}

        members_sql = """
            SELECT username, first_name, last_name, nickname, email, territory
            FROM employees
            WHERE team = :team
            ORDER BY first_name NULLS LAST, last_name NULLS LAST
        """
        rows = conn.execute(text(members_sql), {"team": team_name}).fetchall()

    items = []
    for row in rows:
        full_name = " ".join([p for p in [row[1], row[2]] if p]) or None
        items.append({
            "username": row[0],
            "full_name": full_name,
            "nickname": row[3],
            "email": row[4],
            "territory": row[5]
        })

    return {"team": team_name, "items": items}

@app.get("/")
def serve_index():
    return RedirectResponse(url="/login")

@app.get("/dashboard")
def serve_dashboard(user=Depends(get_current_user)):
    dashboard_path = Path(__file__).parent / "static" / "dashboard.html"
    return FileResponse(dashboard_path)

@app.get("/profile")
def serve_profile(user=Depends(get_current_user)):
    profile_path = Path(__file__).parent / "static" / "profile.html"
    return FileResponse(profile_path)

@app.get("/employees")
def serve_employees(user=Depends(require_admin)):
    employees_path = Path(__file__).parent / "static" / "employees.html"
    return FileResponse(employees_path)

@app.get("/customers")
def serve_customers(user=Depends(require_admin)):
    customers_path = Path(__file__).parent / "static" / "customers.html"
    return FileResponse(customers_path)

@app.get("/customer-summary")
def serve_customer_summary(user=Depends(get_current_user)):
    summary_path = Path(__file__).parent / "static" / "customer_summary.html"
    return FileResponse(summary_path)

@app.get("/update-history")
def serve_update_history(user=Depends(require_admin)):
    history_path = Path(__file__).parent / "static" / "update_history.html"
    return FileResponse(history_path)

@app.get("/help")
def serve_help(user=Depends(get_current_user)):
    help_path = Path(__file__).parent / "static" / "help.html"
    return FileResponse(help_path)

@app.get("/login")
def serve_login():
    login_path = Path(__file__).parent / "static" / "login.html"
    return FileResponse(login_path)

@app.get("/api/me")
def get_me(user=Depends(get_current_user)):
    return {"username": user.get("username"), "role": user.get("role")}

@app.post("/api/login")
async def login(request: Request):
    payload = await request.json()
    username = (payload.get("username") or "").strip().lower()
    password = payload.get("password") or ""

    user = USERS.get(username)
    if user and _verify_password(password, user.get("password_hash", "")):
        request.session["user"] = {"username": username, "role": user.get("role")}
        return {"success": True, "role": user.get("role")}

    emp_sql = text("SELECT username, password_hash FROM employees WHERE username = :username")
    with engine.connect() as conn:
        row = conn.execute(emp_sql, {"username": username}).fetchone()

    if row and row[1] and _verify_password(password, row[1]):
        request.session["user"] = {"username": row[0], "role": "Employee"}
        return {"success": True, "role": "Employee"}

    return JSONResponse({"detail": "invalid_credentials"}, status_code=401)

@app.post("/api/logout")
def logout(request: Request):
    request.session.pop("user", None)
    return {"success": True}

app.mount("/static", StaticFiles(directory="static", html=True), name="static")